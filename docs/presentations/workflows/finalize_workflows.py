from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from pptx import Presentation
import fitz, hashlib, json
from PIL import Image, ImageDraw

root=Path(__file__).resolve().parent
pptx=root/'Portfolio_Intelligence_Demo_Workflows.pptx';pdf=pptx.with_suffix('.pdf')
p=Presentation(pptx);d=fitz.open(pdf);assert len(p.slides)==len(d)==11
qa=root/'qa';qa.mkdir(exist_ok=True)
for i,page in enumerate(d):page.get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False).save(qa/f'slide-{i+1:02}.png')
for start in range(0,len(d),4):
    sheet=Image.new('RGB',(1600,960),'#d4d8de');draw=ImageDraw.Draw(sheet)
    for j in range(start,min(start+4,len(d))):
        im=Image.open(qa/f'slide-{j+1:02}.png');im.thumbnail((780,439));x=(j-start)%2*800;y=(j-start)//2*480
        sheet.paste(im,(x+10,y+25));draw.text((x+12,y+6),f'Slide {j+1}',fill='black')
    sheet.save(qa/f'contact-{start//4+1}.png')
checks=json.loads((root/'powerpoint_layout_check.json').read_text());outside=[]
for i,s in enumerate(p.slides):
    assert s.notes_slide.notes_text_frame.text.strip()
    for sh in s.shapes:
        if sh.left < -10000 or sh.top < -10000 or sh.left+sh.width>p.slide_width+10000 or sh.top+sh.height>p.slide_height+10000:outside.append((i+1,sh.name))
with ZipFile(pptx) as z:
    assert not any(n.startswith('ppt/media/') for n in z.namelist())
    fonts=[n for n in z.namelist() if n.startswith('ppt/fonts/')]
manifest={'slides':11,'editable_shapes':sum(len(s.shapes) for s in p.slides),'raster_images':0,'embedded_fonts':len(fonts),'text_overflows':sum(len(s['text_overflows']) for s in checks),'off_canvas_shapes':outside,'pdf_links':sum(len(page.get_links()) for page in d),'files':[]}
for path in [pptx,pdf]:manifest['files'].append({'name':path.name,'bytes':path.stat().st_size,'sha256':hashlib.sha256(path.read_bytes()).hexdigest()})
(root/'delivery_manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
print(json.dumps(manifest,indent=2))
if manifest['text_overflows']:
    print(json.dumps([s for s in checks if s['text_overflows']],indent=2))
    raise SystemExit('Fix overflow before delivery')
assert not outside
with ZipFile(root/'Demo_Workflow_Package.zip','w',ZIP_DEFLATED) as z:
    for path in [pptx,pdf,root/'README.md',root/'delivery_manifest.json',root/'slide_manifest.json']:z.write(path,path.name)
    folder=root.parent/'client/fonts'
    for name in ['IBMPlexSans-Regular.ttf','IBMPlexSans-Light.ttf','LICENSE.txt']:z.write(folder/name,'fonts/'+name)
print('Package complete')
