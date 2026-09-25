"""Editable, workflow-only demonstration of the implemented dashboard."""
from pathlib import Path
import json, hashlib
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.xmlchemy import OxmlElement
from urllib.parse import quote

ROOT=Path(__file__).resolve().parent
REPO=ROOT.parents[2]
SNAP=REPO/'public/data/api_snapshot.json'
ROUTES=json.loads(SNAP.read_text(encoding='utf-8'))['routes']
ASSET=ROUTES['/portfolio']['assets'][0]
ADV=ROUTES['/assets/DEMO-001/advanced-analysis']
PUBLIC='https://rohen97.github.io/Real-Estate-Portfolio-Multi-Asset-Quant-Model/'
P=Presentation();P.slide_width=Pt(960);P.slide_height=Pt(540)
P.core_properties.title='Portfolio Intelligence | Demo Workflows'
P.core_properties.subject='Overall workflow, selection process and dashboard segments'
P.core_properties.author='Portfolio Intelligence'
BLUE='0F62FE';DARK='202A38';GRAY='657080';LINE='D5DFEA';PALE='EDF5FF'
TEAL='008A80';MINT='E9F6F3';AMBER='A65A00';CREAM='FFF1D8';RED='B52232';PINK='FFF0F1';WHITE='FFFFFF'
FONT='IBM Plex Sans';LIGHT='IBM Plex Sans Light';MANIFEST=[]

def color(c):return RGBColor.from_string(c)
def shape(s,x,y,w,h,fill=WHITE,stroke=LINE,kind=MSO_SHAPE.RECTANGLE,lw=.75):
    z=s.shapes.add_shape(kind,Pt(x),Pt(y),Pt(w),Pt(h))
    z.fill.solid();z.fill.fore_color.rgb=color(fill)
    if stroke:z.line.color.rgb=color(stroke);z.line.width=Pt(lw)
    else:z.line.fill.background()
    if kind==MSO_SHAPE.ROUNDED_RECTANGLE:z.adjustments[0]=.08
    z._element.spPr.append(OxmlElement('a:effectLst'))
    return z
def text(s,x,y,w,h,value,size=14,c=DARK,bold=False,font=FONT,center=False):
    z=s.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h));tf=z.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,ln in enumerate(str(value).split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=ln
        p.alignment=PP_ALIGN.CENTER if center else PP_ALIGN.LEFT;p.space_after=Pt(2);p.line_spacing=1.04
        p.font.name=font;p.font.size=Pt(size);p.font.color.rgb=color(c);p.font.bold=bold
        for r in p.runs:r.font.name=font;r.font.size=Pt(size);r.font.color.rgb=color(c);r.font.bold=bold
    return z
def line(s,x1,y1,x2,y2,c=BLUE,arrow=False,dash=False,lw=1.25):
    z=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(x1),Pt(y1),Pt(x2),Pt(y2))
    z.line.color.rgb=color(c);z.line.width=Pt(lw)
    if dash:z.line.dash_style=MSO_LINE_DASH_STYLE.DASH
    if arrow:
        e=OxmlElement('a:tailEnd');e.set('type','triangle');e.set('w','sm');e.set('len','sm');z.line._get_or_add_ln().append(e)
    return z
def route(s,pts,c=BLUE,dash=False):
    for i,(a,b) in enumerate(zip(pts,pts[1:])):line(s,*a,*b,c,i==len(pts)-2,dash)
def box(s,x,y,w,h,title,body='',fill=PALE,accent=BLUE,ts=16,bs=13):
    shape(s,x,y,w,h,fill,LINE,MSO_SHAPE.ROUNDED_RECTANGLE)
    shape(s,x,y,4,h,accent,None)
    title_h=22*len(title.split('\n'))
    text(s,x+12,y+11,w-23,title_h,title,ts,DARK,True)
    if body:text(s,x+12,y+title_h+17,w-24,h-title_h-21,body,bs,GRAY)
def diamond(s,x,y,w,h,label,fill=CREAM):
    shape(s,x,y,w,h,fill,'FF8B24',MSO_SHAPE.DIAMOND,1.1)
    lines=label.count('\n')+1
    text(s,x+w*.18,y+h/2-lines*8.1,w*.64,lines*18,label,13,AMBER,True,center=True)
def badge(s,x,y,label,w=None,c=BLUE,fill=PALE):
    w=w or len(label)*6.1+20
    shape(s,x,y,w,21,fill,None);text(s,x+8,y+4,w-16,16,label,10.5,c,True)
    return w
def output(s,main,detail=''):
    shape(s,28,464,904,48,PALE,None);shape(s,28,464,4,48,BLUE,None)
    text(s,41,475,70,22,'OUTPUT',11,BLUE,True)
    text(s,117,473,799,24,main,15,BLUE,True)
    if detail:text(s,117,495,799,14,detail,10.3,GRAY)
def inputrow(s,label):
    text(s,28,152,68,20,'INPUT',11,BLUE,True)
    text(s,112,147,811,43,label,15,GRAY)
    line(s,28,198,932,198,LINE,lw=.65)
def slide(title,sub,tabs,notes):
    s=P.slides.add_slide(P.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=color(WHITE)
    text(s,28,17,906,43,title,29,DARK,font=LIGHT)
    text(s,29,64,900,30,sub,13,GRAY)
    line(s,28,99,932,99,BLUE,lw=1)
    text(s,29,112,80,17,'OPEN TABS',9.5,GRAY,True)
    x=113
    for t in tabs:
        w=len(t)*5.8+19
        z=shape(s,x,106,w,22,PALE,None)
        z.click_action.hyperlink.address=PUBLIC+'?tab='+quote(t)
        z=text(s,x+8,111,w-16,16,t,10.3,BLUE,True)
        z.click_action.hyperlink.address=PUBLIC+'?tab='+quote(t)
        x+=w+7
    text(s,29,523,870,13,'Demo overview  |  Fictional assets and saved results  |  25 September 2026',8.5,GRAY)
    text(s,905,521,28,15,f'{len(P.slides):02}',10,GRAY,center=True)
    s.notes_slide.notes_text_frame.text=notes
    MANIFEST.append({'slide':len(P.slides),'title':title,'tabs':tabs,'notes':notes})
    return s

# 1. Overall model: solid core path, separate planning lane into human interpretation.
s=slide('From asset information to a portfolio decision',
        'Read the financial path from left to right; planning studies support a separate analyst review.',
        ['Portfolio','Selection','Optimiser'],
        'Start the demonstration with the overall workflow. The main engine values five owned-asset alternatives, simulates outcomes and ranks expected incremental NPV less 0.3 times nonnegative loss CVaR. Readiness is reported separately and can remain incomplete while demo calculations are shown. Portfolio optimisation uses the full action set, not only each asset winner. The lower planning lane is separate: current ML diagnostics and advanced monthly development values do not feed the core action rankings. The dashed connection represents human interpretation only. The application neither approves nor executes a transaction. Public Pages provides saved results; calculations require the backend. Source: packages/orchestration/engine.py, advanced.py, packages/selection/engine.py and optimisation modules.')
xs=[28,220,412,604,796];w=136
for x,cap in zip(xs,['INPUTS','ANALYSIS','SELECTION','ALLOCATION','HUMAN REVIEW']):text(s,x,161,w,19,cap,10.5,GRAY,True)
for i in range(4):line(s,xs[i]+w,257,xs[i+1],257,arrow=True)
for x,t,b in zip(xs,['Asset + market\ninformation','Cash flows\nand scenarios','Rank actions\n+ show evidence','Optimise actions\nand timing','Review plan\nand open issues'],['Value, income,\nleases, assumptions','Five alternatives;\nvalue and downside','Preferred action;\nreadiness status','Funding, income\nand risk constraints','Accept for further\nunderwriting or revise']):box(s,x,201,w,112,t,b,ts=14.5,bs=11.8,fill=CREAM if x==796 else PALE,accent=AMBER if x==796 else BLUE)
text(s,549,239,54,13,'All actions',9,BLUE,center=True)
for i in range(2):line(s,xs[i]+w,386,xs[i+1],386,TEAL,True)
box(s,28,349,136,75,'Planning inputs','Site, controls,\nlocation context',MINT,TEAL,14.5,11.8)
box(s,220,349,136,75,'Zoning / ML','Capacity screens;\nevidence diagnostics',MINT,TEAL,14.5,11.8)
text(s,360,351,50,27,'Capacity\nonly',9.5,TEAL,center=True)
box(s,412,349,136,75,'Development','Envelope and\nseparate feasibility',MINT,TEAL,14.5,11.8)
route(s,[(548,386),(864,386),(864,313)],TEAL,True)
text(s,596,356,186,41,'Analyst interpretation\nâ€”not an input to ranking',12,TEAL)
output(s,'Ranked asset candidates, evidence gaps and a constrained portfolio plan.',
       'The planning lane is visible context; it does not automatically change the financial ranking.')

# 2. Editable dashboard schematic with actual saved numbers and the actual tab groups.
s=slide('Demo overview | the dashboard follows the decision',
        'A simplified workspace view shows the candidate, its economics and what still needs verification.',
        ['Portfolio','Asset','Selection'],
        'This is an editable presentation schematic, not a screenshot or a new dashboard design. Every navigation group uses existing tab names. DEMO-001 is fictional. Sell has expected incremental NPV 3.756268m and loss CVaR 7.210666m; the risk-adjusted score is 1.593m after the 0.3 penalty. Ranking chart values come from selection.action_comparison. Data Required / Monitor blocks treating the candidate as a verified investment decision. The main public snapshot contains ten fictional assets, while Digital Twins uses a separate synthetic population. Ask the audience to follow one asset from its evidence to its candidate action, then review the portfolio plan.')
shape(s,28,145,904,304,'F8FAFD',LINE)
text(s,42,156,520,24,'Portfolio intelligence  /  DEMO-001',16,DARK,True)
badge(s,778,153,'SAVED DEMO',132)
for x,label,val,col in [(42,'Preferred action','Sell',BLUE),(264,'Expected NPV vs Hold','+S$3.76m',BLUE),(486,'Loss CVaR 95%','S$7.21m',BLUE),(708,'Evidence status','Data Required',AMBER)]:
    shape(s,x,187,209,66,WHITE,LINE);shape(s,x,187,3,66,col,None)
    text(s,x+11,196,188,15,label,10.5,GRAY);text(s,x+11,218,188,27,val,23 if x<708 else 18,col,True)
for x,ww in [(42,270),(324,313),(649,268)]:shape(s,x,265,ww,169,WHITE,LINE)
text(s,54,275,244,23,'01  Open the relevant segment',13.5,DARK,True)
groups=[('Understand','Portfolio Â· Digital Twins'),('Evidence','Asset Â· Evidence'),('Planning','Zoning Â· Zoning ML Â· Development'),('Compare','Valuation Â· Scenario Â· Actions Â· Selection'),('Review','Optimiser Â· Models Â· Architecture Â· Audit')]
for i,(k,v) in enumerate(groups):
    y=305+i*24;text(s,54,y,63,16,k,10,BLUE,True);text(s,119,y,182,23,v,9.5,GRAY)
text(s,336,275,285,23,'02  Compare action scores (S$m)',13.5,DARK,True)
actions=ASSET['selection']['action_comparison'];origin=587;scale=6.3
line(s,origin,302,origin,425,LINE,lw=.75)
for i,a in enumerate(actions):
    y=307+i*23;v=a['risk_adjusted_score_m'];text(s,338,y,76,17,a['action'],11,DARK)
    if v:shape(s,origin+min(v,0)*scale,y+2,abs(v)*scale,10,BLUE if v>0 else 'A7B6C8',None)
    else:shape(s,origin-1,y+2,2,10,TEAL,None)
    text(s,598,y-1,31,18,f'{v:.2f}',9.5,BLUE if v>0 else GRAY)
text(s,661,275,243,23,'03  Explain the decision',13.5,DARK,True)
for i,value in enumerate(['Compare against Hold','Inspect downside risk','Review missing evidence','Test the portfolio fit']):
    y=309+i*27;shape(s,663,y,16,16,BLUE,None,MSO_SHAPE.OVAL);text(s,663,y+1,16,14,str(i+1),9,WHITE,True,center=True);text(s,688,y,215,20,value,12,DARK)
output(s,'A candidate for review, with its rationale and evidence status.',
       'Use the linked tab labels on each slide to jump into the public dashboard.')

# 3. Inputs and readiness.
s=slide('Asset and evidence | establish what is known',
        'Choose one asset, inspect its assumptions and separate supplied information from verified evidence.',
        ['Asset','Evidence'],
        'The Asset page presents the selected asset, economics and provisional recommendation. Evidence shows source details, proxy assumptions and outstanding inputs. The selection readiness function requires current valuation, NOI, site and existing floor area; title, tenure, sources and owner; verified financial and planning evidence; dates within the permitted freshness window; and verified action costs. The diagram summarises analyst review of those checks. Both branches permit a provisional demonstration to be explored; a missing field does not turn an economic recommendation into Hold. Public DEMO-001 lacks verified underwriting evidence.')
inputrow(s,'Selected asset + value and income assumptions + source records, dates and planning context.')
for y,t,b in [(221,'Asset register','Identity, use, location'),(289,'Financial inputs','Value, NOI, leases, costs'),(357,'Evidence records','Title, source, owner, date')]:box(s,28,y,187,56,t,b,ts=14,bs=11.5)
route(s,[(215,249),(235,249),(235,317),(262,317)])
line(s,215,317,262,317,arrow=True);route(s,[(215,385),(235,385),(235,317),(262,317)])
box(s,262,264,207,106,'Check and reconcile','Source, units, missing fields,\nfreshness and verification.',ts=17,bs=13.5)
line(s,469,317,508,317,arrow=True);diamond(s,508,259,125,116,'Evidence\nverified?')
route(s,[(633,317),(671,317),(671,261),(705,261)],TEAL)
route(s,[(633,317),(671,317),(671,377),(705,377)],AMBER)
text(s,651,246,39,17,'YES',10,TEAL,True);text(s,651,393,39,17,'NO',10,AMBER,True)
box(s,705,221,227,79,'Verified input status','Proceed with underwriting review.',MINT,TEAL,16,12.5)
box(s,705,339,227,80,'Provisional screen','Show the missing-evidence list.',CREAM,AMBER,16,12.5)
output(s,'An asset case with visible assumptions and a readiness status.',
       'The demo can continue with provisional inputs; this does not establish decision readiness.')

# 4. Zoning ML is diagnostics, not a forecast-to-trade pipeline.
s=slide('Zoning ML | prioritise the right investigation',
        'The challenger checks consistency; the separate change model is a research capability.',
        ['Zoning ML','Zoning'],
        'Implemented LightGBM code targets current core land use, subtype and GPR band. It includes spatial validation, calibration and abstention machinery. A separate change classifier exists, but a real forecast needs dated change labels, a horizon and calibration evidence. Public outputs are hand-authored fictional fixtures: no eligible trained predictions are displayed. The dashboard suppresses unsupported confidence and change probabilities. The eligibility diamond summarises presentation checks, not a guarantee that an eligible model has investment skill. Current-class discrepancies and future-change scores have distinct targets and are never combined as a trading probability. Neither enters core action rankings. Source: zoning/challenger.py, change_model.py and src/components/zoningEvidence.ts.')
inputrow(s,'Planning records + asset and location features + model training and calibration evidence.')
box(s,28,256,162,110,'Prepare context','Compare source records\nand feature quality.',ts=16,bs=13)
route(s,[(190,311),(215,311),(215,256),(245,256)])
route(s,[(190,311),(215,311),(215,370),(245,370)])
box(s,245,218,205,81,'Current-zoning model','Class / GPR band, uncertainty\nand mismatch checks.',ts=15,bs=12.3)
box(s,245,330,205,81,'Future-change research','Separate change score;\nvalidation required.',MINT,TEAL,15,12.3)
route(s,[(450,259),(479,259),(479,311),(514,311)])
route(s,[(450,370),(479,370),(479,311),(514,311)],TEAL)
diamond(s,514,253,127,117,'Eligible\nestimate?')
route(s,[(641,311),(674,311),(674,253),(706,253)],TEAL)
route(s,[(641,311),(674,311),(674,373),(706,373)],AMBER)
text(s,653,237,40,17,'YES',10,TEAL,True);text(s,654,387,40,17,'NO',10,AMBER,True)
box(s,706,215,226,80,'Analyst investigation','Review a mismatch or\nmonitor a research candidate.',MINT,TEAL,15,12.5)
box(s,706,334,226,80,'Abstain / seek evidence','Do not display unsupported\nprobability estimates.',CREAM,AMBER,15,12.5)
output(s,'A review priority or an explicit lack of evidenceâ€”not a trade instruction.',
       'Public demo: fictional examples; ML outputs are not connected to the main ranking engine.')

# 5. Related capacity and geometry screens, then distinct detailed cases.
s=slide('Zoning and development | test the physical option',
        'Capacity identifies what may fit; a costed development case tests whether it creates value.',
        ['Zoning','Development'],
        'Zoning shows statutory, physical, de-facto and economic GFA, unused capacity and binding controls. Development constructs a synthetic site rectangle and candidate envelopes from related inputs; it does not use a surveyed or approved design. Preferred envelope gross area feeds detailed monthly redevelopment cash flows. En-bloc residual valuation is a related separate calculation using legal GPR and a cost waterfall; it is not a Buy recommendation. Assessed LBC and other verified project evidence are missing in the demo. DEMO-001 has core unused capacity 2184sqm and detailed redevelopment NPV -37.022416m relative to Hold. Its advanced gross area differs from the core adjusted capacity because the assumptions differ. Detailed values are not substituted into main annual rankings. Source: orchestration/advanced.py, actions/detailed_cashflow.py and actions/enbloc.py.')
inputrow(s,'Site area and GPR + height, setbacks and footprint assumptions + construction and project costs.')
box(s,28,258,164,110,'Planning inputs','Use the available controls;\nflag proxies and missing data.',ts=16,bs=12.3)
line(s,192,313,233,313,arrow=True)
box(s,233,239,198,149,'Planning screens','Capacity and unused area\n\nConceptual building envelopes',MINT,TEAL,16,13)
route(s,[(431,313),(457,313),(457,258),(490,258)],TEAL)
route(s,[(431,313),(457,313),(457,375),(490,375)],TEAL)
box(s,490,218,202,85,'Redevelopment case','Monthly costs, income and\nvalue versus Hold.',ts=15.5,bs=12.5)
box(s,490,334,202,85,'En-bloc residual','Land value after development\ncosts and required margin.',ts=15.5,bs=12.5)
route(s,[(692,261),(719,261),(719,313),(751,313)])
route(s,[(692,377),(719,377),(719,313),(751,313)])
box(s,751,250,181,127,'Feasibility review','Incremental NPV\nResidual value\nUnresolved evidence',CREAM,AMBER,16,13)
output(s,'DEMO-001: 2,184 sqm headroom; redevelopment NPV âˆ’S$37.02m versus Hold.',
       'Separate feasibility screen. These detailed values do not feed the main action rankings.')

# 6. Core valuation uses common Hold counterfactual.
s=slide('Valuation and actions | make alternatives comparable',
        'Each action is valued against retaining the same property, using the main annual model.',
        ['Valuation','Actions'],
        'Five implemented owned-asset alternatives are Hold, Retrofit, Repurpose, Redevelop and Sell. The main action engine builds annual cash flows and branches from assumed economics, computes a Hold counterfactual and measures the incremental difference. Project capex, downtime and terminal or disposal value must be consistently included. Its Redevelop template differs from the detailed monthly Development tab and those outputs are not interchangeable. Hold incremental NPV is zero by construction; that does not mean the asset has no absolute risk. Buy is not implemented. Capacity residual, en-bloc and transformation-screen values are excluded from additive NPV adjustments. Source: valuation/engine.py and actions/real_options.py.')
inputrow(s,'Current value and net operating income + rent, vacancy, cap rate, costs and timing assumptions.')
box(s,28,271,175,104,'Hold baseline','Retain income and\nremaining asset value.',ts=17,bs=13.5)
line(s,203,323,239,323,arrow=True)
shape(s,239,217,246,210,'F7F9FC',LINE)
text(s,252,226,220,24,'Build five action cases',16,DARK,True)
for i,(a,b) in enumerate([('Hold','Retain'),('Retrofit','Improve'),('Repurpose','Change use'),('Redevelop','Rebuild'),('Sell','Dispose')]):
    y=258+i*31;shape(s,252,y,220,25,PALE,None);text(s,263,y+5,101,20,a,12.5,BLUE,True);text(s,363,y+5,100,20,b,12.5,GRAY)
line(s,485,323,520,323,arrow=True)
box(s,520,261,186,123,'Value the cash flows','Costs and benefits\nover time, discounted\non a consistent basis.',ts=16,bs=13.5)
line(s,706,323,744,323,arrow=True)
box(s,744,261,188,123,'Incremental value','Action NPV\nminus\nHold NPV',MINT,TEAL,16,14)
output(s,'Five comparable action values, before scenario and portfolio decisions.',
       'Hold is the zero incremental baseline; Buy is not part of the implemented action set.')

# 7. Risk outputs are separate decision measures.
s=slide('Scenario analysis | show the range around the value',
        'Expected value is one outcome measure; downside determines how much uncertainty accompanies it.',
        ['Scenario'],
        'Main simulated actions preserve shared market and approval factors and asset-specific variation. Expected incremental NPV, P10/P50/P90, loss probability and positive-loss CVaR are derived from simulated action outcomes. CVaR95 is the average in the worst five-percent loss tail; it is not a percentile value or a probability. The Selection page subtracts 0.3 times nonnegative loss CVaR from mean incremental NPV. DEMO-001 Sell: expected 3.756268m; probability of loss .2444; CVaR 7.210666m. These are fictional assumed scenarios, not observed investment performance. The scenario page follows the selected action and does not choose the maximum raw NPV independently. Source: scenarios/engine.py and selection/engine.py.')
inputrow(s,'Action cash flows + shared market, approval and cost assumptions + asset-specific uncertainty.')
box(s,28,270,174,109,'Set uncertainty','Shared factors and\nasset-specific variation.',ts=16,bs=13.2)
line(s,202,324,242,324,arrow=True)
box(s,242,270,174,109,'Simulate outcomes','Generate a distribution\nfor every action.',ts=16,bs=13.2)
for mid in [244,325,405]:route(s,[(416,324),(442,324),(442,mid),(474,mid)])
for y,t,b in [(214,'Expected value','Average incremental NPV'),(295,'Outcome range','P10 / P50 / P90'),(375,'Downside','Loss probability and loss CVaR')]:box(s,474,y,235,59,t,b,ts=14.5,bs=11.5)
for mid in [244,325,405]:route(s,[(709,mid),(732,mid),(732,324),(759,324)])
box(s,759,256,173,135,'Read the trade-off','Does the expected gain\njustify the modelled\ndownside?',CREAM,AMBER,16,14)
output(s,'DEMO-001 Sell: +S$3.76m expected NPV; 24.4% loss probability; S$7.21m CVaR.',
       'CVaR describes the average loss in the worst 5% of simulated outcomes.')

# 8. Selection ranking then explicit readiness diamond and human follow-through.
s=slide('Selection | rank the action, then test readiness',
        'The preferred financial action and the operational evidence status answer different questions.',
        ['Selection','Actions'],
        'The implemented score is expected incremental NPV minus 0.3 times max(0, loss CVaR95). Rank all five alternatives, then require asset evidence and verified action costs/cash flows before describing the candidate as ready for portfolio review. The no branch changes the management signal to Data Required / Monitor but retains the preferred screening action. DEMO-001 Sell mean NPV 3.756268 less .3*7.210666 gives 1.593068m, versus Hold score zero. The model therefore favours Sell financially but reports Data Required / Monitor. This is not an executable sell instruction. The yes branch is a candidate for further human and portfolio review, not automatic investment approval. Buy absent; no demonstrated alpha. Source: packages/selection/engine.py.')
inputrow(s,'Hold Â· Retrofit Â· Repurpose Â· Redevelop Â· Sell, each with expected NPV and downside estimates.')
box(s,28,265,204,124,'Risk-adjusted ranking','Expected NPV\nâˆ’ 0.3 Ã— loss CVaR\n\nCompare with Hold.',ts=16,bs=13.5)
line(s,232,327,268,327,arrow=True)
box(s,268,277,181,101,'Preferred action','Keep the economic\ncandidate visible.',ts=16,bs=13.5)
line(s,449,327,487,327,arrow=True);diamond(s,487,270,133,114,'Evidence\nready?')
route(s,[(620,327),(665,327),(665,255),(706,255)],TEAL)
route(s,[(620,327),(665,327),(665,387),(706,387)],AMBER)
text(s,642,240,40,18,'YES',10,TEAL,True);text(s,642,402,40,18,'NO',10,AMBER,True)
box(s,706,217,226,79,'Candidate for review','Check the portfolio fit\nand investment rationale.',MINT,TEAL,15.5,12.5)
box(s,706,348,226,79,'Data Required / Monitor','Resolve the evidence gaps\nbefore acting.',CREAM,AMBER,15.5,12.5)
output(s,'DEMO-001: Sell scores +S$1.59m; the evidence status remains Data Required.',
       'Data Required / Monitor does not mean that Hold wins the economic comparison.')

# 9. Portfolio optimisation consumes the full action set, with constraint audit.
s=slide('Optimiser | choose a plan that fits the portfolio',
        'An attractive standalone action may change once funding, retained income and execution limits apply.',
        ['Optimiser'],
        'Input all action alternatives for the portfolio, not just independently selected asset winners. Main optimisation chooses actions and start years while checking capex, cumulative liquidity, retained income, project concurrency and development share. Joint loss CVaR uses aligned scenarios rather than summing marginal risks. The returned schedule is rechecked by the constraint audit. It remains a conditional model plan with synthetic inputs. The two-stage calculation creates contingent plans across scenarios; stability reruns perturb assumptions. Public Pages displays saved runs and disables new backend calculation controls. Funding released by sale is not profit, and annual liquidity buckets permit within-year netting. Source: optimisation/multi_period.py, two_stage.py, constraints.py and the Optimiser UI.')
inputrow(s,'All asset alternatives + capital and liquidity budgets + income, project and development limits.')
box(s,28,264,188,111,'Portfolio alternatives','Value and risk for\neach action and asset.',ts=16,bs=13)
line(s,216,319,252,319,arrow=True)
box(s,252,264,192,111,'Choose action + year','Optimise value subject\nto shared constraints\nand joint downside.',ts=16,bs=13)
line(s,444,319,482,319,arrow=True);diamond(s,482,259,130,120,'Constraint\nchecks\npass?')
route(s,[(612,319),(654,319),(654,253),(700,253)],TEAL)
route(s,[(612,319),(654,319),(654,379),(700,379)],AMBER)
text(s,630,237,44,17,'YES',10,TEAL,True);text(s,630,395,44,17,'NO',10,AMBER,True)
box(s,700,215,232,80,'Allocation + annual plan','Actions, start years, capital,\nremaining liquidity and income.',MINT,TEAL,15.5,12.3)
box(s,700,338,232,80,'Revise and rerun','Review the failed constraint\nand adjust the proposed plan.',CREAM,AMBER,15.5,12.3)
text(s,29,422,610,27,'Other views: two-stage contingent plans and assumption-sensitivity checks.',11.5,GRAY)
output(s,'A feasible model schedule with explicit constraints and funding requirements.',
       'Public demo displays saved plans; new calculations require the backend.')

# 10. Digital twins are an environment comparison, not a historical backtest.
s=slide('Digital twins | compare decisions across environments',
        'Explore a separate synthetic population and see which choices change when assumptions move.',
        ['Digital Twins'],
        'Digital Twins presents separate generated histories, lease and action scenarios, and a synthetic population. Environment selection loads the saved results for that environment. Independent preferred actions and the full constrained portfolio allocation are distinct paths. Asset/search filters change visible rows and asset plots but do not rerun or subset the full portfolio optimisation. Changing decisions is a sensitivity flag. The example rows on this slide describe display fields and do not invent asset outcomes. Simulation regret, action match and interval coverage are generator diagnostics, not historical backtests or realised alpha. Source: src/components/DigitalTwinDashboard.tsx and digital-twin snapshot methodology.')
inputrow(s,'Choose an environment, then filter the synthetic assets whose decisions you want to inspect.')
box(s,28,266,172,100,'Select environment','Load the saved\nassumption case.',ts=16,bs=13)
route(s,[(200,316),(228,316),(228,250),(262,250)])
route(s,[(200,316),(228,316),(228,377),(262,377)])
box(s,262,210,246,81,'Independent asset choices','Preferred action, NPV range\nand probability of loss.',ts=15.5,bs=13)
box(s,262,337,246,81,'Full portfolio allocation','Constrained action and start year\nfor the whole population.',ts=15.5,bs=13)
route(s,[(508,251),(547,251),(547,316),(587,316)])
route(s,[(508,378),(547,378),(547,316),(587,316)])
box(s,587,251,345,131,'Compare what changes','Which assets change action?\nWhere does the portfolio choose differently?\nWhich assumptions drive the difference?',MINT,TEAL,17,13.5)
output(s,'A sensitivity view of asset choices, portfolio choices and their differences.',
       'Filtering the asset table does not rerun the full portfolio plan. Results are synthetic experiments.')

# 11. Supporting segments end in human interpretation rather than automated approvals.
s=slide('Review segments | explain the result before acting',
        'Use diagnostics, the model map and the evidence trail to explain the recommendation and its limits.',
        ['Models','Architecture','Audit'],
        'Models provides forecast and diagnostic views, but public forecast cards are formula fixtures, not a fitted production forecast. Architecture explains the implemented calculation modules. Audit surfaces evidence and readiness checks; it is not a complete immutable transaction approval log. The diagram is a user review workflow: these three screens do not automatically combine into a recommendation service. A human investment decision and any implementation approval occur outside the demonstration. End the demo by naming the candidate, its downside, the binding portfolio constraint if any, and the missing information that would change the decision. No investment alpha has been established.')
inputrow(s,'The candidate action and portfolio plan + diagnostic results + assumptions and source information.')
for y,t,b in [(215,'Models','Inspect diagnostics and data status.'),(296,'Architecture','Understand the calculation path.'),(377,'Audit','Review input readiness and gaps.')]:box(s,28,y,271,59,t,b,ts=14.5,bs=12)
for mid in [245,326,407]:route(s,[(299,mid),(330,mid),(330,326),(367,326)])
box(s,367,256,250,140,'Explain the decision','Why this candidate?\nWhat could make it wrong?\nWhich inputs need verification?\nWhat changes the portfolio fit?',ts=17,bs=13.5)
route(s,[(617,326),(666,326),(666,326),(708,326)],AMBER,True)
box(s,708,264,224,123,'Human investment review','Proceed to underwriting,\nrequest more evidence,\nor revise the proposal.',CREAM,AMBER,16,13.5)
output(s,'An understandable decision brief and a clear next review step.',
       'The dashboard supports analysis; investment approval and execution happen outside the demo.')

assert len(P.slides)==11
TARGET=ROOT/'Portfolio_Intelligence_Demo_Workflows.pptx'
P.save(TARGET)
with ZipFile(TARGET) as z:contents={n:z.read(n) for n in z.namelist()}
for n,raw in list(contents.items()):
    if n.startswith('ppt/theme/') and n.endswith('.xml'):
        tree=etree.fromstring(raw);ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
        for e in tree.xpath('//a:fontScheme/a:majorFont/a:latin',namespaces=ns):e.set('typeface',LIGHT)
        for e in tree.xpath('//a:fontScheme/a:minorFont/a:latin',namespaces=ns):e.set('typeface',FONT)
        for e in tree.xpath('//a:effectStyle',namespaces=ns):
            for child in list(e):e.remove(child)
            e.append(OxmlElement('a:effectLst'))
        contents[n]=etree.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
with ZipFile(TARGET,'w',ZIP_DEFLATED) as z:
    for n,raw in contents.items():z.writestr(n,raw)
(ROOT/'slide_manifest.json').write_text(json.dumps({'slides':MANIFEST,'source_snapshot_sha256':hashlib.sha256(SNAP.read_bytes()).hexdigest(),'reference':'Four user-provided process and demo-overview screenshots; original editable diagrams.'},indent=2),encoding='utf-8')
print(json.dumps({'file':str(TARGET),'slides':len(P.slides),'editable_shapes':sum(len(s.shapes) for s in P.slides)},indent=2))
