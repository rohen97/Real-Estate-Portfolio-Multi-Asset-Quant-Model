"""Rebuild the formal Word interpretation from the final captured public API data."""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime, timezone
from hashlib import sha256
import json
from pathlib import Path
import shutil
import subprocess
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs/reports"
PUBLIC = ROOT / "public/reports"
FILENAME = "Real_Estate_Dashboard_Interpretation_Report.docx"


def read(path):
    return json.loads(path.read_text(encoding="utf-8"))


def number(value, places=2):
    return "Unavailable" if value is None else f"{value:,.{places}f}"


def percent(value, places=1):
    return "Unavailable" if value is None else f"{value * 100:.{places}f}%"


def field(paragraph, instruction):
    run = paragraph.add_run()
    for tag, kind in (("begin", "w:fldChar"), (instruction, "w:instrText"), ("separate", "w:fldChar"), ("end", "w:fldChar")):
        element = OxmlElement(kind)
        if kind == "w:fldChar":
            element.set(qn("w:fldCharType"), tag)
        else:
            element.set(qn("xml:space"), "preserve")
            element.text = f" {tag} "
        run._r.append(element)


def build(skip_charts=False):
    OUT.mkdir(parents=True, exist_ok=True)
    PUBLIC.mkdir(parents=True, exist_ok=True)
    if not skip_charts:
        subprocess.run([sys.executable, str(Path(__file__).with_name("make_charts.py"))], cwd=ROOT, check=True)
    snapshot = read(ROOT / "public/data/api_snapshot.json")
    routes = snapshot["routes"]
    core = routes["/portfolio"]["assets"]
    twins = routes["/digital-twins/dashboard"]
    assets, environments = twins["assets"], twins["environments"]
    opt, two, stability = [routes[path] for path in ("/snapshots/optimise", "/snapshots/two-stage", "/snapshots/stability")]
    diagnostic = routes["/models/backtest"]
    market = routes["/market/public"]
    revision = routes["/model/revision"]
    manifest_path = OUT / "chart_manifest.json"
    if not manifest_path.exists():
        manifest_path = OUT / "figures/chart_manifest.json"
    manifest = read(manifest_path)
    figures = manifest if isinstance(manifest, list) else manifest.get("figures", manifest.get("charts", []))
    if isinstance(manifest, dict) and manifest.get("sources"):
        for key, source in (("api_snapshot_sha256", ROOT / "public/data/api_snapshot.json"),
                            ("digital_twin_dashboard_sha256", ROOT / "public/data/digital_twin_dashboard.json")):
            if manifest["sources"].get(key) != sha256(source.read_bytes()).hexdigest():
                raise ValueError("Charts were generated from an older snapshot; rerun without --skip-charts")
    if len(figures) < 20:
        raise ValueError("The detailed report requires at least twenty generated figures")
    if any(asset.get("asset_id", "").startswith("DEMO-") is False for asset in core) or len(core) != 10:
        raise ValueError("This report is scoped to the ten explicitly fictional public assets")

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.2677), Inches(11.6929)
    section.top_margin = section.bottom_margin = Inches(.64)
    section.left_margin = section.right_margin = Inches(.75)
    section.header_distance, section.footer_distance = Inches(.23), Inches(.28)
    for name in ("Normal", "Body Text"):
        style = doc.styles[name]
        style.font.name, style.font.size = "Calibri", Pt(10.5)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.10
    for name, size, color in (("Title", 31, "123C54"), ("Subtitle", 15, "557183"), ("Heading 1", 19, "123C54"), ("Heading 2", 13.5, "176B95"), ("Heading 3", 11, "176B95")):
        style = doc.styles[name]
        style.font.name, style.font.size = "Calibri", Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_after = Pt(10)
    doc.styles["Caption"].font.size = Pt(8.5)
    doc.styles["Caption"].font.color.rgb = RGBColor.from_string("405D70")
    header = section.header.paragraphs[0]
    header.text = "REAL ESTATE PORTFOLIO INTELLIGENCE  |  MODEL AND DASHBOARD INTERPRETATION"
    header.style = "Caption"
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Fictional assets · Official macro context · Page ").font.size = Pt(8)
    field(footer, "PAGE")
    footer.add_run(" of ").font.size = Pt(8)
    field(footer, "NUMPAGES")
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings.element.append(update)
    doc.core_properties.title = "Real Estate Portfolio Intelligence: Dashboard and Model Interpretation"
    doc.core_properties.author = "Prepared for Rohen"
    doc.core_properties.subject = "Revised model, official market context, fictional portfolio and complete dashboard interpretation"

    def p(text, lead=None, style=None):
        paragraph = doc.add_paragraph(style=style)
        if lead:
            paragraph.add_run(lead + " ").bold = True
        paragraph.add_run(str(text))
        return paragraph

    def heading(text, level=2):
        return doc.add_heading(text, level)

    def page(title):
        doc.add_page_break()
        heading(title, 1)

    def table(headers, rows, widths=None):
        result = doc.add_table(rows=1, cols=len(headers))
        result.style = "Light Shading Accent 1"
        result.autofit = False
        for index, (cell, value) in enumerate(zip(result.rows[0].cells, headers)):
            cell.text = str(value)
            if widths:
                cell.width = Inches(widths[index])
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        result.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            for index, (cell, value) in enumerate(zip(result.add_row().cells, row)):
                cell.text = str(value)
                if widths:
                    cell.width = Inches(widths[index])
        for row in result.rows:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(4)
                    paragraph.paragraph_format.space_before = Pt(3)
                    paragraph.paragraph_format.line_spacing = 1.03
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
        p("")
        return result

    total_value = sum(asset["economics"]["current_value_m"] for asset in core)
    total_noi = sum(asset["economics"]["current_noi_m"] for asset in core)
    twin_value = sum(asset["valuation_m"] for asset in assets)
    twin_noi = sum(asset["noi_m"] for asset in assets)
    captured = snapshot["metadata"]["generated_at"][:10]
    p("FORMAL DASHBOARD AND MODEL REPORT", style="Subtitle")
    p("")
    doc.add_heading("Real Estate Portfolio\nIntelligence Dashboard", 0)
    p("Revised decision logic, financial methods, data provenance, graphs and interpretation", style="Subtitle")
    p("")
    table(["Report basis", "Details"], [
        ["Prepared for", "Rohen"], ["Snapshot date", captured], ["Model", revision.get("model_version")],
        ["Coverage", f"All 15 dashboard tabs; {len(figures)} explanatory figures"],
        ["Asset population", "10 fictional Singapore demonstration assets"],
        ["Economic evidence", "Official aggregate market context; synthetic/proxy asset underwriting"],
        ["Publication", "All-tab saved public dashboard plus a separately deployable calculation API"],
        ["Report version", "2.0 — revised cashflows, risk, scheduling and evidence controls"],
    ], [1.55, 5.1])
    p("The revised model provides a more consistent comparison of retaining, improving and selling assets. Its numerical results remain conditional on fictional asset data and unvalidated action assumptions. No graph in this report demonstrates realised investment performance or proven market alpha.", lead="Principal conclusion")
    p("All amounts are in Singapore dollars. “m” means S$ million. Current value, cash proceeds, annual NOI, incremental NPV, CVaR and annual return measure different things and must not be added or compared as interchangeable quantities.")

    page("Contents and reading guide")
    field(doc.add_paragraph(), 'TOC \\o "1-1" \\h \\z \\u')
    p("Word refreshes this contents field when the document is opened or its fields are updated. The report first establishes the evidence and decision context, then explains each figure, and finally records financial formulas, optimisation assumptions, governance, sources and reproducibility.")
    table(["Reader's objective", "Recommended sequence"], [
        ["Understand the current conclusion", "Executive interpretation → portfolio and environment figures → selection/readiness"],
        ["Understand one asset", "Asset/Evidence → Valuation/Actions/Scenario → Zoning/Development → Audit"],
        ["Understand allocation", "Standalone recommendations → portfolio constraints → joint risk → recourse → stability"],
        ["Assess the evidence", "Data provenance → Models/forecast diagnostics → limitations → source register"],
    ], [2.0, 4.65])
    p("Each figure is reconstructed from the same saved JSON used by the public dashboard. These are publication figures, not screenshots. Their source, units, interpretation and limitations are given immediately beside the figure. A figure can consolidate more than one dashboard panel or add a reconciliation to make the underlying calculation easier to assess.")

    page("1. Executive interpretation")
    core_mix = Counter(asset["recommendation"]["action"] for asset in core)
    opt_mix = Counter(row["action"] for row in opt["selections"])
    core_mix_text = " and ".join(f"{count} {action}" for action, count in core_mix.items())
    opt_mix_text = " and ".join(f"{count} {action}" for action, count in opt_mix.items())
    p(f"The core portfolio contains S${number(total_value)}m of proxy value and S${number(total_noi)}m of annual proxy NOI. Its independent action ranking produces {core_mix_text}. Every asset remains Data Required / Monitor because numerical ranking does not supply verified underwriting, planning evidence or observed investment outcomes.")
    p(f"Under the default portfolio constraints, the multi-period model selects {opt_mix_text}. Expected incremental NPV is S${number(opt['portfolio_expected_npv_m'])}m; the positive-loss CVaR is S${number(opt['portfolio_cvar_95_m'])}m and the penalised objective is S${number(opt['risk_adjusted_objective_m'])}m. The cash released of S${number(opt.get('capital_released_m'))}m is a separate funding measure, not the incremental benefit.")
    table(["Twin portfolio measure", "Base", "Stress", "Structural"], [
        ["Expected incremental NPV, S$m", *[number(env["portfolio"]["portfolio_expected_npv_m"]) for env in environments]],
        ["Committed construction, S$m", *[number(env["portfolio"]["capital_required_m"]) for env in environments]],
        ["Sale cash released, S$m", *[number(env["portfolio"]["capital_released_m"]) for env in environments]],
        ["Audited constraint violations", *[env["portfolio"]["constraint_violations"] for env in environments]],
        ["Synthetic interval coverage", *[percent(env["p10_p90_coverage"]) for env in environments]],
    ], [2.65, 1.33, 1.33, 1.34])
    p(f"The Digital Twins path independently generates S${number(twin_value)}m of value and S${number(twin_noi)}m of annual NOI. Its levels differ from the core proxies because the generator uses different areas, rents, expenses, yields and sector profiles. The difference is not wealth created by the new model, and the two portfolios must not be summed.")
    p("The absence of selected construction projects in a particular solution is an economically possible response to the assumed capex, disruption, approval and required return. It is not evidence that redevelopment cannot work in Singapore. Equally, a Sell ranking is conditional on an assumed sale price relative to retained income; it is not a transaction recommendation for an identified real building.", lead="Decision interpretation")

    page("2. Evidence, provenance and data improvement")
    table(["Evidence layer", "Current content", "Interpretation"], [
        ["Public market context", f"Official Singapore sector indices; snapshot as of {market.get('as_of', 'unavailable')}", "Aggregate price/rent context with source and vintage metadata"],
        ["Asset identities and locations", "10 DEMO assets and illustrative coordinates", "Fictional, not a verified ownership or holdings list"],
        ["Financial underwriting", "Proxy core values and generated twin values/histories", "0% independently verified underwriting"],
        ["Zoning and site geometry", "Fixture labels and synthetic rectangles", "No statutory entitlement or title-boundary conclusion"],
        ["Core risk and twin outcomes", "Model-generated common scenarios and separate synthetic evaluation", "Assumption sensitivity, not realised outcomes"],
        ["Forecast diagnostics", "Arithmetic examples whose metrics reconcile to rows", "No trained public forecast model or real historical accuracy"],
    ], [1.55, 2.35, 2.75])
    p(f"The former context-completeness label has been separated from verified provenance. There are {twins['portfolio_summary']['real_portfolio_anchors']} verified real portfolio anchors and {twins['portfolio_summary'].get('fictional_asset_contexts', 10)} fictional asset contexts. Average populated-context coverage is {percent(twins['portfolio_summary']['average_real_context_pct'])}; it does not measure confidence, accuracy or verified information.")
    p("Source categories now map explicitly: Office to Commercial; Logistics/Warehouse/Industrial to Industrial; Hospitality to Hotel; Mixed use to a distinct Mixed Use profile. Unknown sectors become Other. This changes underwriting assumptions as well as labels and removes the earlier silent Residential fallback.")
    p("Official quarterly index observations improve the source quality of displayed market context. They are neither property appraisals nor an estimated alpha signal. Sector mapping, observation date, retrieval date, staleness and publication/revision limitations travel with the inputs. A current downloaded historical series does not establish what an investor knew at a past decision date.")

    page("3. Complete dashboard guide: portfolio and asset views")
    table(["Tab", "Purpose and controls", "Interpretation rule"], [
        ["Portfolio", "Search, country scope, open an asset; value/action/source summaries", "Scope controls affect the visible population; provisional action does not override readiness"],
        ["Digital Twins", "Environment selector and asset/action filters; common scenario comparisons", "Separate independently preferred actions from the constrained portfolio selections"],
        ["Asset", "Selected property's identity, economics, action and status", "Check asset ID and financial basis before comparing another page"],
        ["Evidence", "Source records, missing fields, proxy inputs and verification requirements", "Populated fields and source-record counts do not establish verification"],
        ["Zoning", "Map radius and polygon labels; capacity waterfall/ranges and sensitivity", "Fixture map geometry is illustrative; capacity is a screening calculation"],
        ["Zoning ML", "Label comparison, discrepancy, illustrative confidence/change displays", "No public trained classifier; fixture agreement is not classification accuracy"],
        ["Development", "3D envelope, view context, en-bloc residual and monthly cashflows", "Read feasibility and missing-cost blockers before any return number"],
        ["Selection", "Ranking, hurdle, diagnostic overlays and readiness gate", "An internal preferred action can coexist with Data Required / Monitor"],
    ], [1.05, 2.75, 2.85])
    p("Opening a portfolio row sets the shared selected asset. Asset, Evidence, Zoning, Development, Selection, Valuation, Scenario, Actions and Audit follow that selection. List filters should not be confused with recalculating an individual property. Plotly hover and zoom expose numeric chart details; Word figures are static snapshots of those quantities.")

    page("3.1 Complete dashboard guide: model and portfolio decisions")
    table(["Tab", "Purpose and controls", "Interpretation rule"], [
        ["Valuation", "Annual NOI/terminal cashflow and present-value comparison", "Terminal cash is already inside DCF; current value and DCF are alternative benchmarks"],
        ["Scenario", "Stored action draws, quantiles, loss probability, factor correlation and approval stages", "P10–P90 is a central outcome interval, not an error bar on mean NPV"],
        ["Actions", "Five action means, downside, capital and execution periods", "Negative sale capex is cash released; it does not mean construction subsidy"],
        ["Optimiser", "Multi-period allocation, two-stage recourse and assumption perturbation", "Read saved/live mode, request settings, risk definition and constraint audit"],
        ["Models", "Forecast diagnostics and market-regime availability", "Unavailable training or calibration is shown as unavailable, not as measured zero skill"],
        ["Architecture", "Evidence-to-decision, asset, portfolio and training/governance flows", "A diagrammed component is not proof that its private inputs or trained artifact exist"],
        ["Audit", "Model/source identifiers, verification gates and review requirements", "Gate codes are categorical status, not percentage confidence or approval likelihood"],
    ], [1.05, 2.75, 2.85])
    p("The public GitHub Pages build now exposes all tabs using saved API-shaped results. Its saved optimiser panels reflect the explicit default requests recorded in the snapshot. A browser filter can explore the saved data, but cannot invoke new Python calculations on GitHub Pages.")
    p("A hosted or locally running Python backend supports the calculation endpoints. The interface must identify that mode before enabling a new optimisation run. The static transport rejects POST calls and excludes private decision/review records. This report documents the snapshot and deployable backend capability; it does not assert that an unverified external service URL is live.")

    for index, entry in enumerate(figures, 1):
        title = entry.get("title", f"Figure {index}")
        page(f"Figure {index}. {title}")
        section_name = entry.get("section", "Dashboard interpretation")
        p(str(section_name), style="Subtitle")
        file = Path(entry["file"])
        candidates = [file, ROOT / file, OUT / file, OUT / "figures" / file.name]
        image_path = next((candidate for candidate in candidates if candidate.exists()), None)
        if image_path is None:
            raise FileNotFoundError(entry["file"])
        image_paragraph = doc.add_paragraph()
        image_paragraph.paragraph_format.keep_with_next = True
        image_paragraph.add_run().add_picture(str(image_path), width=Inches(6.65))
        p(f"Units: {entry.get('units', 'As labelled')}. Source: {entry.get('source', 'Captured public API snapshot')}. Reconstructed from saved data.", style="Caption")
        interpretation = entry.get("interpretation", "")
        if isinstance(interpretation, list):
            for part in interpretation:
                p(part)
        else:
            for part in str(interpretation).split("\n\n"):
                p(part)
        p(entry.get("limitation", "Fictional asset inputs; interpretation is conditional on stated assumptions."), lead="Interpretation boundary")

    page("4. Financial method: a common Hold counterfactual")
    p("The annual model values unlevered asset cashflows. NOI in future year t grows from current NOI at the assumed annual growth rate. Terminal value capitalises the following year's NOI at the exit cap rate, less selling costs. All future income is discounted at the end of its year; committed action capex is a time-zero outflow.")
    table(["Quantity", "Definition"], [
        ["Annual NOI", "NOI(t) = NOI(0) × (1 + growth)^t"],
        ["Terminal value", "TV(T) = NOI(T + 1) ÷ exit cap rate × (1 − selling-cost rate)"],
        ["Retained-asset PV", "Σ NOI(t)/(1 + discount rate)^t + TV(T)/(1 + discount rate)^T"],
        ["Action incremental NPV", "Expected action project PV − retained-asset PV"],
        ["Branch expectation", "Approval probability × success NPV + failure probability × failure NPV"],
        ["Sell", "Immediate net sale proceeds − retained-asset PV"],
        ["Hold", "Exactly zero incremental NPV and capex; absolute retained-asset PV remains positive"],
    ], [1.5, 5.15])
    p("NOI uplift starts after execution; construction reduces income during execution, including partial years. A failed development retains Hold income and loses 55% of planned capex in the annual screening engine. That recovery assumption is explicit and uncalibrated. It is a different, simplified failure mechanism from the advanced monthly engine's predevelopment sunk fees.")
    p("Full committed costs remain in the capital budget even when probability-weighted expected expenditure is smaller. Funding a successful project requires its actual committed cash. Financing interest is excluded from unlevered action NPV; equity cashflows require a separate financing ledger and equity discount rate.")
    p("Default terminal uplift is zero because NOI uplift is already capitalised. Residual planning capacity and legacy Wait/Phase/Abandon/Expand scores do not receive an unsupported additional option premium. Zero credited premium means not valued, not an estimate of zero market optionality.")

    page("5. Core risk: common scenarios and exact loss tails")
    p("The core engine generates 5,000 aligned factor scenarios for rent growth, vacancy, cap rates, interest rates, construction costs and approval delay. The factor hierarchy uses a calibrated regime artifact if available, calibrated covariance next, then explicit fallback assumptions. In the public run, missing calibration is visible; downloading price/rent indices alone does not create a fully calibrated six-factor process.")
    p("Fallback draws combine correlated Gaussian and variance-standardised correlated heavy-tail shocks. The mixture is variance-normalised before rate/delay bounds. The displayed heatmap is the realised correlation of the bounded model sample, with the input correlation matrix retained separately. Neither matrix demonstrates an observed economic relationship in the present fallback run.")
    p("Action shocks scale with asset value, capex, annual NOI and execution duration. Sell reverses exposure to the income forgone by sale. Approval samples select explicit success/failure cashflows; failed approval retains Hold economics and only cost uncertainty remains. The direct borrowing-rate penalty has been removed from unlevered NPV. These local sensitivity coefficients remain assumptions.")
    table(["Risk statistic", "Meaning"], [
        ["Expected NPV", "Mean of the generated incremental outcomes"],
        ["P10 / P50 / P90", "10th, 50th and 90th outcome quantiles; a central 80% range"],
        ["Probability of loss", "Fraction of incremental NPV draws below zero"],
        ["Positive-part loss", "L = max(0, −incremental NPV)"],
        ["VaR95", "95th percentile of positive-part loss"],
        ["CVaR95", "Average loss in the worst 5% probability mass, splitting a boundary observation where needed"],
    ], [1.5, 5.15])
    p("A rare loss cannot be offset by profits merely because the 5% tail also contains profitable observations. For example, a 2% probability of losing S$100m and otherwise gaining S$100m gives positive-loss CVaR95 of S$40m. Signed tail statistics are stored separately. Exact saved samples feed the distribution chart and the cards, so those displays now reconcile.")

    page("6. Digital Twins: conditional environments and independent evaluation")
    p("The twin generator produces its own underwriting and twelve years of financial history. These histories are synthetic, including apparent cycles, tenant diversification and expense trends. Predictions use observable generated fields, not hidden simulation truth. Small bounded score adjustments remain screening assumptions: maximum absolute valuation adjustments are 1.5% for Retrofit and 2% for Repurpose/Redevelop, applied only after success.")
    table(["Assumption", "Base", "Stress", "Structural"], [
        ["Headline rent-level change", *[percent(env["assumptions"]["rent_change"]) for env in environments]],
        ["Cap-rate shift, percentage points", *[number(env["assumptions"]["cap_rate_shift"] * 100) for env in environments]],
        ["Construction-cost multiplier", *[number(env["assumptions"]["construction_cost_multiplier"]) for env in environments]],
        ["Approval multiplier", *[number(env["assumptions"]["approval_multiplier"]) for env in environments]],
    ], [2.65, 1.33, 1.33, 1.34])
    p("Segment demand adds to the headline rent-level change. Both Hold and action cashflows use that revised NOI. Annual growth is separately bounded at −3% to 8%; exit cap rates at 2% to 15%; sale-price multipliers at 0.35 to 1.75. Construction inflation changes both funding and NPV, and sale-price stress changes both sale proceeds and incremental value. NOI-continuity constraints use the environment-specific baseline.")
    p("Each action has 2,048 predictive draws with shared Student-t(6) economic/construction factors and idiosyncratic shocks. A separate seed generates one synthetic evaluation outcome, with hidden effects available only to that evaluation. The hindsight oracle observes those outcomes; the predictive rule does not. Oracle agreement and regret therefore assess a closed simulation, not actual management skill.")
    p(f"Nominal 80% interval coverage is {[percent(env['p10_p90_coverage']) for env in environments]} across the three cases, each using 50 asset/action evaluation observations. The shortfall is visible evidence that these assumed predictive ranges are not calibrated to every hidden evaluation effect. Recommendation stability across three cases is different from the optimiser's twenty-run perturbation stability.")

    page("7. Portfolio optimisation: timing, cash and joint downside")
    p("The multi-period mixed-integer model selects exactly one action/start option per asset. It maximises expected incremental NPV minus a chosen penalty times aggregate positive-loss CVaR. When aligned action scenarios are supplied, they retain the same scenario index across the portfolio; the optimiser aggregates them before calculating risk. Without aligned scenarios it explicitly falls back to summed marginal risk penalties.")
    p("The solver uses 128 aligned scenarios for computational tractability; core risk cards use 5,000 and twin cards 2,048. A joint-versus-marginal diversification comparison must use the same 128 observations and weights. The optimiser now reports that comparable marginal sum separately from the larger-sample input-card sum.")
    table(["Constraint or convention", "Interpretation"], [
        ["Annual and cumulative capital", "Nominal spending minus receipts; the reserve must survive every year"],
        ["NOI retention", "Retained baseline NOI at or above the selected minimum each year"],
        ["Concurrency and development share", "Active projects and exact floor(asset count × share) limits; zero means zero"],
        ["Complete project horizon", "All construction/funding must fit; unfinished costs cannot be truncated"],
        ["Delayed execution", "NPV at start is discounted to the decision date; passive pre-start carry contributes zero incremental NPV"],
        ["Sale settlement", "Receipts fund spending only after settlement; income is retained before settlement"],
        ["Independent audit", "Recompute every selected-schedule constraint; tolerance 0.000001; failed audit prevents feasible status"],
    ], [1.65, 5.0])
    p("Budget and liquidity numbers are not additional investment profits. Negative net funding is a receipt surplus. The annual cash convention does not resolve the timing of receipts and payments within the same year. Delaying an aggregate NPV is also an approximation: it does not re-underwrite future market rents, lease expiries, cost inflation or terminal dates.")
    p("A feasible optimal plan is best under the encoded objective and assumptions. It does not certify legal capacity, affordability for a specific company, lender covenants or real forecast accuracy. The retained financing constraint based on a percentage of capex is a financing proxy, not balance-sheet loan-to-value.")

    page("8. Two-stage recourse and stability")
    p("The two-stage model commits to an initial plan and permits specified actions to execute or cancel after one of three illustrative scenarios resolves. Probabilities describe those assumed scenarios; they are not market-implied or empirically estimated frequencies. The default scenario system is coarser than the 128-draw joint-risk optimiser, so their CVaR values are not interchangeable.")
    p("Cancellation uses the asset's actual Hold economics and Hold funding plus a default 8% cancellation fee. Without a feasible Hold alternative, cancellation is unavailable. This prevents a fictitious cheap cancellation from replacing an expensive baseline. Scenarios are assumed known before construction: the model does not learn continuously through the life of a project.")
    p("Scenario delays move capex, disruption, contractor workload and completion. Each scenario and year enforces income continuity, concurrency, annual funding and cumulative liquidity. Future sales cannot fund earlier spending. Dependency checks require prerequisite completion before dependent execution. Sale-price multipliers affect cash receipts and economic NPV together; the downside default is 0.90 and the upside default 1.05.")
    p("For negative-NPV actions, a downside multiplier worsens the loss rather than making it smaller. Increased construction spending also reduces NPV rather than only making the budget harder to satisfy. The top-level annual plan is probability-weighted recourse; each scenario retains its actual plan and audit.")
    p(f"The saved stability calculation has {stability.get('successful_runs')} successful and {stability.get('failed_runs')} failed perturbation runs, with {stability.get('fragile_assets')} fragile assets under its threshold. It perturbs model assumptions around the baseline and reports modal-action frequencies only among feasible audited runs. Failure frequency is reported separately.")
    p("A collection of per-asset modal actions need not be jointly feasible. A 60% frequency is not a 60% probability that an investment is correct. Perturbation objective quantiles are sensitivity summaries, not prediction intervals or out-of-sample investment returns. The stability report separates risk-adjusted objective from unpenalised expected NPV.")

    page("9. Planning, capacity and evidence limits")
    p("Statutory screening GFA equals assumed site area multiplied by assumed plot ratio. Physical capacity is the lesser of statutory GFA and the geometric envelope implied by footprint, coverage, setbacks and height. De-facto capacity applies a further practical factor; expected-approved capacity applies an assumed approval probability. These are successive definitions, not quantities to sum.")
    p("Economic capacity depends on development value less all-in construction allowances. Unused economic GFA and residual value are screening diagnostics. Economic GFA may exceed probability-weighted expected-approved GFA because one represents an economic feasibility calculation and the other weights physical space by approval assumptions. Neither proves a legally available entitlement.")
    p("Sensitivity bars show the absolute response to the stated finite changes in GPR, coverage, height, setbacks or approval probability. They are not unit-free elasticities. Zero height sensitivity can mean that plot ratio is already binding, rather than that height regulations never matter. Capacity P10/P90 values reflect generated uncertainty, not surveyed boundaries.")
    p("The approval path contains conditional stages. Overall pass probability multiplies conditional probabilities; percentages cannot be added. Simulated durations can include paths that stop early on failure, so unconditional timing percentiles should not be presented as completion time conditional on approval.")
    p("The current map polygons, zoning labels and site rectangles are fictional. Radius controls change the displayed area, not confidence in title boundaries or planning controls. A detailed-looking envelope remains a screening geometry. If no feasible envelope exists, redevelopment value is unavailable rather than calculated using current floor area as a substitute entitlement.")
    p("Missing surrounding-building geometry now produces unavailable view metrics. An empty obstruction chart does not establish an unobstructed view or absence of future supply. Fixture legal/predicted labels may agree by construction; the challenger abstains and identifies those numbers as hand-authored display examples, without trained-model or statutory authority.")

    page("10. Advanced monthly development and financing")
    p("The advanced redevelopment model uses monthly construction and financing cashflows. Construction costs apply to all rebuilt GFA, including replacement of the existing building. Additional GFA is retained only for separate capacity/LBC screening. Charging only incremental area would omit the cost of rebuilding existing floors.")
    p("The project includes the current asset's opportunity value at time zero. Its Hold comparison includes the same owned asset and horizon. For owned-property development, incremental NPV equals project NPV minus Hold NPV, so the common opportunity cost cancels in the comparison while remaining in total project returns. Acquire is a distinct comparison with making no acquisition and includes applicable buyer duty.")
    table(["Monthly identity", "Treatment"], [
        ["Unlevered project cashflow", "Operating/terminal cash minus development spending and relevant transaction costs"],
        ["Equity cashflow", "Project cashflow + debt draws − cash interest − principal repayment"],
        ["Debt balance", "Prior debt + draws − repayments; development debt repaid at disposal"],
        ["Incremental NPV", "Expected project NPV − Hold NPV on the same timing basis"],
        ["Equity NPV", "Unavailable unless an explicit equity discount rate is supplied"],
        ["Equity IRR", "Includes asset opportunity cost; unavailable for nonconventional/ambiguous cashflows"],
    ], [1.65, 5.0])
    p("Predevelopment professional fees are at risk before approval. A failed branch retains the asset's Hold income and value while losing those fees; a successful branch proceeds through construction, lease-up and sale. Probabilities weight complete branches. Higher leverage cannot improve unlevered project NPV merely by adding debt cash to asset proceeds.")
    p("The current public advanced outputs remain incomplete: LBC is unassessed, tax-use allocation and development costs are proxies, and existing debt, acquisition financing, corporate income tax, ABSD and tenure top-up are not fully provided. Numeric compatibility fields explicitly exclude missing costs. They are conditional screening values, not executable transaction appraisals.")

    page("11. En-bloc residual and decision readiness")
    p("The en-bloc screen starts with completed-development value and deducts construction, professional/finance allowances, marketing, demolition, tenant costs, known charges and developer margin. The offered land value must also bear buyer duty calculated on that offer. The solver reconciles offer plus duty to the residual available before duty.")
    p("A missing LBC remains null and a decision blocker. A supplied rate is still a screening input rather than a verified statutory assessment. An externally assessed zero is different from a missing charge. Actual chargeable uplift, sector, use, tenure and applicable schedules require verification.")
    p("The consent screen can use supplied observed support, or an explicitly uncalibrated independent equal-owner binomial approximation when real voting shares are missing. Raising the required consent threshold cannot increase success probability. This is not a legal voting determination; actual ownership shares and relevant legal thresholds remain external evidence requirements.")
    p("The old success-weighted sale value is now identified as only one component. If a sale attempt fails, the existing asset is retained. The incremental comparison includes this retained-asset fallback and compares it with Hold at the same future transaction date. Common intervening income cancels only under the stated assumption that the attempt does not disrupt it.")
    p("Residual value is not a cash receipt guaranteed to the owner, an annual return, or an asset appraisal. A negative premium versus the current assumed value can make an en-bloc attempt unattractive even when the gross completed-development value is large. The waterfall is informative only if every cost category and missing charge is visible.")
    p("Selection now ranks the common risk-adjusted incremental action comparison. Unsupported transformation, view, capacity and en-bloc premiums remain diagnostics rather than arbitrary additions to alpha. Readiness checks for finite, verified financials, current planning evidence and model inputs govern whether a decision is supported. Public fictional assets remain blocked from transaction-ready status.")

    page("12. Forecast validation, possible alpha and what is not established")
    p("The forecasting implementation now constructs adjacent fiscal-year samples for the same asset, rejects duplicate years and invalid features, and keeps features and labels aligned. Missing covariates are not silently set to zero. A minimum amount of temporal data is required before training proceeds.")
    p("Earlier years train the standardised Ridge, gradient-boosted point and quantile models; the penultimate target year calibrates intervals; the latest target year remains a final holdout. Scaling and fitting see training rows only. The fitted artifact is not subsequently refitted on test outcomes, and its training/calibration/test cutoffs are stored.")
    p("The held-out report compares MAE and RMSE with a no-change prior-NOI forecast, records relative forecast skill and evaluates interval width/coverage/score. Negative skill means the fitted ensemble underperformed the simple baseline. Calibration on time-dependent panel data does not inherit an unconditional independent-observation coverage guarantee, and fiscal-year order does not establish publication availability at an investment date.")
    p(f"The public dashboard does not contain that trained-model evaluation. Its {diagnostic['noi_forecast']['observations']} displayed fictional formula rows have MAE S${number(diagnostic['noi_forecast']['mae_m'], 4)}m, bias S${number(diagnostic['noi_forecast']['bias_m'], 4)}m and {percent(diagnostic['noi_forecast']['p10_p90_coverage'])} interval coverage, all recomputed from those exact rows. Their apparent fit is an arithmetic demonstration, not market forecasting evidence.")
    p("Action hit rates remain unavailable when independent action outcomes and valid counterfactual labels do not exist. An observed management action is not necessarily the economically optimal action; measuring returns only when model and management agree introduces selection bias and leaves alternatives unobserved.")
    p("Potential research improvements include verified lease-level NOI, contemporaneous capex/approval data, better timing of expenditure and sale receipts, explicit opportunity costs, and out-of-time performance measurement against simple alternatives. Establishing investment alpha additionally requires point-in-time information, transaction costs, implementability, risk exposures and a defensible counterfactual. The current work improves the measurement system but does not supply that missing evidence.")

    page("13. Architecture, hosting and governance")
    p("The architecture follows evidence ingestion → validated asset context → zoning/financial assumptions → cashflow and uncertainty models → action ranking → constrained allocation → review. The dashboard presents those layers; it does not convert unverified source fields into facts. Separate core, twin and advanced paths remain explicitly identified because their underwriting and counterfactuals differ.")
    p("The separate lease engine can use tenant-level area, passing/market rent, collection and expiry dates to build monthly income. It prorates the original tenancy, applies renewal or one reletting gap at the supplied expiry, and avoids repeating that gap in every later year. Subsequent expiries, incentives, breaks and re-letting costs still require explicit schedules. The current public dashboard has no verified company leases; this capability is not evidence that its proxy values were derived from a real rent roll.")
    p("GitHub Pages serves the compiled React application and saved JSON. This release includes a public API-shaped snapshot for all tabs, allowing asset navigation, graph exploration, search and saved optimiser results on any computer without running Python. The public snapshot excludes review/decision records and hidden latent simulation truth.")
    p("New calculations require the Python backend. Build-time preparation generates fictional model data before service startup, and the hosted calculation endpoints validate bounded numeric inputs. A configured frontend API URL connects that service to the same interface. A deployment configuration is not evidence that an authenticated external hosting account has completed its deployment.")
    p("The Audit tab communicates source, model, recommendation and human-review gates. A coloured gate is a categorical status; it is not a percentage complete or probability of success. The source code hash and market-snapshot hash identify the captured calculation. Any subsequent package/data change should regenerate the results and the report before making a comparison.")
    p("Solver audit verifies the encoded schedule; data gates verify whether required evidence is present; human review determines whether the evidence and assumptions are acceptable. These are separate responsibilities. Missing underwriting, title evidence, LBC or forecast calibration remains visible rather than being replaced by optimistic zeroes or apparently precise certainty scores.")

    page("14. Revision summary and verification evidence")
    table(["Earlier problem", "Revised behaviour", "Why it matters"], [
        ["Year-one cash discounted at time zero", "Common year-end DCF and exact zero incremental Hold", "Stops timing and baseline differences manufacturing value"],
        ["Duplicate capacity/terminal premiums", "Capacity unpriced unless justified; no unsupported option credit", "Avoids overlapping benefits counted twice"],
        ["Only extra rebuilt area charged", "Construction charged on full replacement GFA", "Represents the economic cost of rebuilding"],
        ["Latent simulation truth in prediction", "Observable predictor and disjoint synthetic evaluation", "Separates decision information from hindsight"],
        ["Fixed currency risk and resampled chart", "Size-scaled, shared scenarios and exact stored chart samples", "Improves exposure consistency and reconciliation"],
        ["Marginal risk called portfolio risk", "Joint positive-loss CVaR on aligned scenarios", "Reflects modeled dependence and coherent tail penalties"],
        ["Truncated projects and loose recourse", "Whole schedules, per-year constraints and independent audits", "Prevents unfinanced or temporally impossible selections"],
        ["Fixed metrics and fictional real anchors", "Recomputed metrics, explicit no-training status and provenance", "Prevents unsupported accuracy and data-quality claims"],
    ], [2.0, 2.45, 2.2])
    validation_path = OUT / "validation_results.json"
    if validation_path.exists():
        validation = read(validation_path)
        table(["Verification", "Recorded result"], [[key.replace("_", " "), "; ".join(f"{field.replace('_', ' ')}: {value}" for field, value in values.items()) if isinstance(values, dict) else str(values)] for key, values in validation.items()], [1.5, 5.15])
    else:
        p("Focused regression suites cover valuation identities, cash ledgers, approval branches, empirical risk reconciliation, forecasting leakage, schedule feasibility and source status. Final repository-wide execution results are recorded separately in validation_results.json when available; this report does not invent a final test count.")
    p("Passing tests demonstrates the specified software invariants and exercised paths. It does not verify real asset inputs, legal entitlement, economic calibration or future investment performance. The new tests include rare-loss CVaR, failed-approval exposure, exact forecast-summary reconciliation and a final-year leakage perturbation.")

    page("15. Asset register and current decision outputs")
    table(["Asset", "Source sector", "Value S$m", "NOI S$m", "Core rank", "Plan"], [
        [asset["asset_id"], ", ".join(asset.get("segments", [])), number(asset["economics"]["current_value_m"]),
         number(asset["economics"]["current_noi_m"]), asset["recommendation"]["action"],
         next(row["action"] for row in opt["selections"] if row["asset_id"] == asset["asset_id"])] for asset in core
    ], [1.0, 1.4, .9, .9, 1.1, 1.35])
    p("Value and NOI in this table are the core model's proxy economics. The portfolio plan column is constrained and may differ from the independent core rank. The Digital Twins table uses a different generated underwriting dataset and is intentionally not reconciled to these monetary levels. All ten asset records are fictional and all remain subject to data verification.")
    p("A large sale receipt can support liquidity while the incremental gain remains small: retaining the asset already has substantial value. Similarly, high gross development value can coexist with negative redevelopment NPV after the existing asset, replacement construction, downtime and required return are recognised. These distinctions explain much of the revised dashboard behaviour.")

    page("16. Sources, model files and reproducibility")
    sources = []
    for segment, context in market.get("market_indicators", {}).items():
        for series, metadata in context.get("series", {}).items():
            sources.append([segment.replace("_", " ") + " " + series, metadata.get("agency", "Official source"), metadata.get("source_url", "Unavailable"), metadata.get("period_end", "Unavailable")])
    table(["Market series", "Agency", "Source URL", "Period end"], sources, [1.2, 1.0, 3.35, 1.1])
    p("Official aggregate indices are downloaded reference observations. Source date, period end and refresh date are separate. History may be revised; this report does not claim point-in-time availability for trading or causally attribute asset performance to a public index movement.")
    p("Primary method references: NYU Stern / Aswath Damodaran valuation framework (https://pages.stern.nyu.edu/~adamodar/New_Home_Page/lectures/val.html); Singapore Land Authority Land Betterment Charge (https://www.sla.gov.sg/properties/land-betterment-charge/); scikit-learn chronological evaluation documentation (https://scikit-learn.org/stable/modules/generated/sklearn.model_selection.TimeSeriesSplit.html); Romano, Patterson and Candès, Conformalized Quantile Regression, NeurIPS 2019 (https://proceedings.neurips.cc/paper/2019/hash/5103c3584b063c431bd1268e9b5e76fb-Abstract.html). These references explain methodology and do not verify the fictional assets.")

    page("16.1 Reproduction record and remaining limitations")
    table(["Artifact", "Purpose"], [
        ["public/data/api_snapshot.json", "Captured all-tab route data, exact chart samples and saved optimiser requests"],
        ["public/data/digital_twin_dashboard.json", "Public twin result snapshot and explicit methodology/provenance"],
        ["data/examples/demo_full_model_results.json", "Canonical core action results and underlying scenario observations"],
        ["scripts/create_public_demo.py", "Reproducible fictional inputs and diagnostic fixtures"],
        ["scripts/build_public_api_snapshot.py", "Allowlisted export; rejects non-demo portfolios and private source inputs"],
        ["scripts/report/make_charts.py", "Recreates every figure from saved data"],
        ["scripts/report/build_report.py", "Recreates this Word report and source/hash manifest"],
        ["docs/*_revision.md", "Detailed valuation, optimisation, simulation and forecasting changes"],
    ], [3.0, 3.65])
    p(f"Captured model version: {revision.get('model_version')}. Package code SHA-256: {revision.get('code_sha256')}. Market snapshot SHA-256: {revision.get('market_snapshot_sha256')}. The source manifest beside this report hashes the exact data, chart manifest and method documents used for reproduction.")
    p("Remaining limits include fictional portfolio economics; uncalibrated action uplifts, approvals and sensitivities; proxy site geometry and tax-use allocation; missing assessed LBC and other project costs; coarse recourse scenarios; 128-scenario portfolio tail sampling; annual cash timing; incomplete borrower/debt covenants; unavailable disclosure vintages and counterfactual policy outcomes. These are evidence and modeling limits, not facts that can be removed by additional chart detail.")
    p("The practical use is to inspect assumptions, compare consistent cashflow alternatives, understand risk and funding, and identify what evidence a decision requires. Any real asset decision needs verified underwriting and current professional planning, valuation, tax and transaction input.")

    target = OUT / FILENAME
    doc.save(target)
    shutil.copy2(target, PUBLIC / FILENAME)
    source_paths = [ROOT / "public/data/api_snapshot.json", ROOT / "public/data/digital_twin_dashboard.json", manifest_path,
                    ROOT / "docs/valuation_revision.md", ROOT / "docs/optimisation_revision.md", ROOT / "docs/simulation_revision.md", ROOT / "docs/forecasting_revision.md",
                    Path(__file__), Path(__file__).with_name("make_charts.py")]
    if validation_path.exists():
        source_paths.append(validation_path)
    source_manifest = {"report_version": "2.0", "generated_at": datetime.now(timezone.utc).isoformat(),
                       "model_revision": revision, "figures": len(figures),
                       "sources": [{"file": path.relative_to(ROOT).as_posix(), "sha256": sha256(path.read_bytes()).hexdigest(), "bytes": path.stat().st_size} for path in source_paths],
                       "report": {"file": target.relative_to(ROOT).as_posix(), "sha256": sha256(target.read_bytes()).hexdigest(), "bytes": target.stat().st_size}}
    (OUT / "source_manifest.json").write_text(json.dumps(source_manifest, indent=2), encoding="utf-8", newline="\n")
    print(json.dumps({"docx": str(target), "public_docx": str(PUBLIC / FILENAME), "figures": len(figures), "paragraphs": len(doc.paragraphs), "tables": len(doc.tables)}, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--skip-charts", action="store_true")
    build(parser.parse_args().skip_charts)
