from __future__ import annotations
from collections import defaultdict
from hashlib import sha1
from pathlib import Path
import json,re,time
import httpx
from openpyxl import load_workbook

def text(value): return str(value).strip() if value is not None else ''
def key_name(value): return re.sub(r'[^a-z0-9]+','',text(value).lower())
def parse_number(value):
 if value in (None,''): return None
 m=re.search(r'[0-9][0-9,]*(?:\.[0-9]+)?',text(value));return float(m.group().replace(',','')) if m else None

def country_for(sheet,first):
 if sheet=='For Rent (SG)': return 'Singapore'
 if sheet=='Hotels & Serviced Res': return first or 'Unknown'
 if first.startswith('Australia'): return 'Australia'
 return 'Singapore'

def load_source_workbook(path:Path)->list[dict]:
 wb=load_workbook(path,read_only=True,data_only=True);rows=[]
 for ws in wb.worksheets:
  headers=[text(x.value) for x in next(ws.iter_rows(min_row=1,max_row=1))]
  for index,values in enumerate(ws.iter_rows(min_row=2,values_only=True),start=2):
   if not any(values): continue
   record=dict(zip(headers,[text(v) for v in values]));first=text(values[0]);
   if ws.title=='For Rent (SG)':
    segment,property_name,address,size,status=[text(v) for v in values[:5]];detail=status
   elif ws.title=='Sale, Malls, AU, Storage':
    segment,property_name,address,detail=[text(v) for v in values[:4]];size=''
   else:
    country,segment,property_name,operator=[text(v) for v in values[:4]];address='';detail=operator;first=country;size=''
   if not property_name: continue
   rows.append(dict(source_sheet=ws.title,source_row=index,country=country_for(ws.title,first),segment=segment,property_name=property_name,address=address,details=detail,size_from_sqft=parse_number(size),asking_rent_monthly_sgd=parse_number(detail) if 'month' in detail.lower() else None,raw=record))
 return rows

def infer_tenure(items):
 combined=' '.join(x.get('details','') for x in items).lower()
 if 'freehold' in combined:return {'type':'Freehold','lease_years':None,'commencement_year':None,'remaining_years':None,'source':'portfolio catalogue detail'}
 match=re.search(r'(\d{2,3})\s*(?:yrs|years)',combined);start=re.search(r'(?:from|commencing)\s*(19\d{2}|20\d{2})',combined)
 if match:
  years=int(match.group(1));year=int(start.group(1)) if start else None;remaining=max(0,years-(2026-year)) if year else None;return {'type':f'{years}-year leasehold','lease_years':years,'commencement_year':year,'remaining_years':remaining,'source':'portfolio catalogue detail'}
 return {'type':'Unknown','lease_years':None,'commencement_year':None,'remaining_years':None,'source':'not supplied'}
def merge_assets(rows:list[dict])->list[dict]:
 grouped=defaultdict(list)
 for row in rows: grouped[(row['country'],key_name(row['property_name']))].append(row)
 assets=[]
 for (_,name_key),items in grouped.items():
  primary=max(items,key=lambda x:(bool(x['address']),len(x['details'])))
  asset_id='FEO-'+sha1((primary['country']+name_key).encode()).hexdigest()[:8].upper()
  segments=sorted(set(x['segment'] for x in items if x['segment']))
  addresses=sorted(set(x['address'] for x in items if x['address']))
  sizes=[x['size_from_sqft'] for x in items if x['size_from_sqft']]
  rents=[x['asking_rent_monthly_sgd'] for x in items if x['asking_rent_monthly_sgd']];tenure=infer_tenure(items);details=sorted(set(x['details'] for x in items if x['details']))
  assets.append(dict(asset_id=asset_id,name=primary['property_name'],country=primary['country'],jurisdiction='Singapore URA' if primary['country']=='Singapore' else 'Outside Singapore scope',segments=segments,address=addresses[0] if addresses else '',source_addresses=addresses,source_records=len(items),size_from_sqft=min(sizes) if sizes else None,asking_rent_monthly_sgd=min(rents) if rents else None,source_sheets=sorted(set(x['source_sheet'] for x in items)),catalogue_details=details,tenure_hint=tenure,source_rows=[dict(sheet=x['source_sheet'],row=x['source_row']) for x in items],observed_financials=False,synthetic_financials=True))
 return sorted(assets,key=lambda x:(x['country']!='Singapore',x['name']))

def geocode_singapore(assets:list[dict],cache_path:Path,delay=1.25)->list[dict]:
 cache=json.loads(cache_path.read_text()) if cache_path.exists() else {};client=httpx.Client(timeout=30,headers={'User-Agent':'Far-East-Real-Estate-POC/0.2'})
 try:
  for asset in assets:
   if asset['country']!='Singapore': continue
   query=asset['address'] or asset['name'];postal=re.search(r'(?<!\d)(\d{6})(?!\d)',query);query=postal.group(1) if postal else query
   cached=cache.get(query,'missing');retry=cached=='missing' or cached is None or (isinstance(cached,dict) and '429' in cached.get('error',''))
   if retry:
    result=None;last_error=None
    variants=[]
    for candidate in [query,asset['address'].replace(', Singapore','').replace('Singapore','').strip(),asset['address'].split('/')[0].split(',')[0].strip(),asset['name']]:
     if candidate and candidate not in variants:variants.append(candidate)
    for candidate in variants:
     for attempt in range(6):
      try:
       response=client.get('https://www.onemap.gov.sg/api/common/elastic/search',params={'searchVal':candidate,'returnGeom':'Y','getAddrDetails':'Y','pageNum':1})
       if response.status_code==429:
        last_error=f'429 retry {attempt+1}';time.sleep(max(delay,2**attempt));continue
       response.raise_for_status();items=response.json().get('results',[]);result=items[0] if items else None;last_error=None;break
      except Exception as exc:last_error=str(exc);time.sleep(max(delay,2**attempt))
     if result:break
    cache[query]=result if not last_error else {'error':last_error}
    cache_path.parent.mkdir(parents=True,exist_ok=True);cache_path.write_text(json.dumps(cache,indent=2),encoding='utf-8');time.sleep(delay)
   result=cache.get(query)
   if result and not result.get('error'):
    asset.update(latitude=float(result['LATITUDE']),longitude=float(result['LONGITUDE']),postal_code=result.get('POSTAL',''),geocoded_address=result.get('ADDRESS',''),geocode_source='Singapore OneMap API')
   else: asset.update(geocode_review_required=True)
 finally: client.close()
 return assets

def build_portfolio(workbook_path:Path,output_path:Path,geocode=True,overrides_path:Path|None=None)->list[dict]:
 assets=merge_assets(load_source_workbook(workbook_path))
 if overrides_path and overrides_path.exists():
  overrides=json.loads(overrides_path.read_text())
  for asset in assets:
   override=overrides.get(key_name(asset['name']))
   if override:
    asset['address']=override['address'];asset['address_override_source']=override.get('source');asset['address_override_verified_date']=override.get('verified_date')
 if geocode: assets=geocode_singapore(assets,output_path.parent/'onemap_geocode_cache.json')
 output_path.parent.mkdir(parents=True,exist_ok=True);output_path.write_text(json.dumps(assets,indent=2),encoding='utf-8');return assets
