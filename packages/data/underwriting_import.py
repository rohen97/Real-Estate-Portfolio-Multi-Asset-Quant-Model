from __future__ import annotations
from datetime import date,datetime
from pathlib import Path
import json,math
from openpyxl import load_workbook
CURRENT_DATE=date(2026,9,23)
CORE_FIELDS=['ownership_entity','ownership_pct','title_lot','tenure','site_area_sqm','existing_gfa_sqm','nla_sqm','occupancy','current_valuation_m','valuation_date','noi_m','current_use','data_owner','evidence_source','evidence_date','verification_status']
def norm(value):
 if isinstance(value,str):return value.strip()
 return value
def number(value):
 if value in (None,''):return None
 try:return float(value)
 except:return None
def iso(value):
 if value in (None,''):return None
 if isinstance(value,(datetime,date)):return value.date().isoformat() if isinstance(value,datetime) else value.isoformat()
 try:return datetime.fromisoformat(str(value)).date().isoformat()
 except:return None
def row_map(headers,values):return {str(h).strip():norm(v) for h,v in zip(headers,values)}
def validate_asset_row(row):
 mapped={'asset_id':row.get('Asset ID'),'ownership_entity':row.get('Ownership Entity'),'ownership_pct':number(row.get('Ownership %')),'title_lot':row.get('Title Lot / MK-TS'),'tenure':row.get('Tenure'),'tenure_start':iso(row.get('Tenure Start')),'tenure_expiry':iso(row.get('Tenure Expiry')),'site_area_sqm':number(row.get('Site Area sqm')),'existing_gfa_sqm':number(row.get('Existing GFA sqm')),'nla_sqm':number(row.get('NLA sqm')),'occupancy':number(row.get('Occupancy %')),'current_valuation_m':number(row.get('Current Valuation SGD m')),'valuation_date':iso(row.get('Valuation Date')),'noi_m':number(row.get('NOI SGD m')),'passing_rent_m':number(row.get('Passing Rent SGD m pa')),'opex_m':number(row.get('Operating Expenses SGD m pa')),'debt_balance_m':number(row.get('Debt Balance SGD m')),'interest_rate':number(row.get('Interest Rate %')),'loan_maturity':iso(row.get('Loan Maturity')),'current_use':row.get('Current Use'),'data_owner':row.get('Data Owner'),'evidence_source':row.get('Evidence Source'),'evidence_date':iso(row.get('Evidence Date')),'verification_status':row.get('Verification Status')}
 errors=[];warnings=[]
 for field in CORE_FIELDS:
  if mapped.get(field) in (None,''):errors.append({'field':field,'code':'required','message':'Required field is missing'})
 if mapped['ownership_pct'] is not None and not 0<mapped['ownership_pct']<=1:errors.append({'field':'ownership_pct','code':'range','message':'Ownership must be greater than 0 and no more than 100%'})
 for field in ['site_area_sqm','existing_gfa_sqm','nla_sqm','current_valuation_m']:
  if mapped[field] is not None and mapped[field]<=0:errors.append({'field':field,'code':'positive','message':'Value must be greater than zero'})
 if mapped['existing_gfa_sqm'] and mapped['nla_sqm'] and mapped['nla_sqm']>mapped['existing_gfa_sqm']:errors.append({'field':'nla_sqm','code':'relationship','message':'NLA cannot exceed existing GFA'})
 if mapped['occupancy'] is not None and not 0<=mapped['occupancy']<=1:errors.append({'field':'occupancy','code':'range','message':'Occupancy must be between 0% and 100%'})
 if mapped['noi_m'] is not None and mapped['passing_rent_m'] is not None and mapped['opex_m'] is not None:
  expected=mapped['passing_rent_m']-mapped['opex_m']
  if abs(expected-mapped['noi_m'])>max(.1,abs(mapped['noi_m'])*.05):warnings.append({'field':'noi_m','code':'reconciliation','message':'NOI does not reconcile to passing rent less operating expenses within 5%'})
 if mapped['verification_status']!='Verified':errors.append({'field':'verification_status','code':'verification','message':'Record must be marked Verified before import'})
 return mapped,errors,warnings
def parse_sheet(ws):
 headers=[cell.value for cell in ws[4]];rows=[]
 for values in ws.iter_rows(min_row=5,values_only=True):
  if not any(v not in (None,'') for v in values):continue
  rows.append(row_map(headers,values))
 return rows
def import_underwriting(workbook_path:Path,base_portfolio_path:Path,output_dir:Path):
 wb=load_workbook(workbook_path,data_only=True,read_only=True);portfolio=json.loads(base_portfolio_path.read_text());by_id={x['asset_id']:x for x in portfolio};report={'workbook':str(workbook_path),'imported_at':datetime.now().isoformat(),'assets_total':len(portfolio),'assets_ready':0,'errors':[],'warnings':[],'sheet_counts':{}}
 observed={};asset_rows=parse_sheet(wb['Asset Register']);report['sheet_counts']['Asset Register']=len(asset_rows)
 for row in asset_rows:
  aid=row.get('Asset ID');mapped,errors,warnings=validate_asset_row(row)
  if aid not in by_id:errors.append({'field':'asset_id','code':'unknown','message':'Asset ID is not in the base portfolio'})
  if errors:report['errors'].append({'asset_id':aid,'sheet':'Asset Register','issues':errors})
  else:observed[aid]=mapped;report['assets_ready']+=1
  if warnings:report['warnings'].append({'asset_id':aid,'sheet':'Asset Register','issues':warnings})
 zoning_rows={r.get('Asset ID'):r for r in parse_sheet(wb['Zoning & Planning'])};report['sheet_counts']['Zoning & Planning']=len(zoning_rows)
 for aid,data in observed.items():
  z=zoning_rows.get(aid,{})
  data['zoning_verified']=z.get('Title Boundary Verified')=='Yes' and z.get('Current Use Permitted')=='Yes'
  data['height_limit_m']=number(z.get('Height Limit m'));data['site_coverage_max']=number(z.get('Site Coverage %'));data['front_setback_m']=number(z.get('Front Setback m'));data['rear_setback_m']=number(z.get('Rear Setback m'));data['side_setback_m']=number(z.get('Side Setback m'));data['written_permission_ref']=z.get('Written Permission Ref');data['planning_reviewer']=z.get('Planning Reviewer');data['title_boundary_geojson_path']=z.get('Title Boundary GeoJSON Path')
  if not data['zoning_verified']:report['warnings'].append({'asset_id':aid,'sheet':'Zoning & Planning','issues':[{'field':'title/current use','code':'planning_verification','message':'Observed financial inputs imported, but planning capacity remains blocked until title boundary and current use are both verified'}]})
 for sheet in ['Financials','Operational & ESG','Leases','Debt','Capex','Planning History','Legacy Scorecard Inputs','Transactions & Comps']:
  rows=parse_sheet(wb[sheet]);report['sheet_counts'][sheet]=len(rows)
  output_dir.mkdir(parents=True,exist_ok=True);(output_dir/(sheet.lower().replace(' & ','_and_').replace(' ','_')+'.json')).write_text(json.dumps(rows,indent=2,default=str),encoding='utf-8')
 enriched=[]
 for asset in portfolio:
  copy=dict(asset)
  if asset['asset_id'] in observed:copy['observed_inputs']=observed[asset['asset_id']
  ]
  enriched.append(copy)
 output_dir.mkdir(parents=True,exist_ok=True);(output_dir/'portfolio_enriched.json').write_text(json.dumps(enriched,indent=2,default=str),encoding='utf-8');(output_dir/'underwriting_validation_report.json').write_text(json.dumps(report,indent=2,default=str),encoding='utf-8')
 return report
